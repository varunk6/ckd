import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Constants
COLOR_ORANGE = RGBColor(0xFF, 0x6B, 0x00)
COLOR_ORANGE_LIGHT = RGBColor(0xFF, 0xF0, 0xE5)
COLOR_DARK = RGBColor(0x11, 0x18, 0x27)
COLOR_MUTED = RGBColor(0x4B, 0x55, 0x63)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_BG_CARD = RGBColor(0xF9, 0xFA, 0xFB)
COLOR_BORDER = RGBColor(0xE5, 0xE7, 0xEB)
COLOR_GREEN = RGBColor(0x10, 0xB9, 0x81)
COLOR_RED = RGBColor(0xEF, 0x44, 0x44)

blank_layout = prs.slide_layouts[6]

def set_white_bg(slide):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE

def add_header(slide, tag, title, subtitle):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.text = tag.upper()
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_ORANGE

    tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.5))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = title
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = COLOR_DARK

    tb3 = slide.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(11.7), Inches(0.4))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = subtitle
    p3.font.size = Pt(12)
    p3.font.color.rgb = COLOR_MUTED

def add_card(slide, left, top, width, height, title, items, border_color=COLOR_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_BG_CARD
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)

    tb = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.15), Inches(width - 0.3), Inches(height - 0.3))
    tf = tb.text_frame
    tf.word_wrap = True

    if title:
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_DARK

    if items:
        for idx, item in enumerate(items):
            p = tf.add_paragraph() if (title or idx > 0) else tf.paragraphs[0]
            p.text = f"• {item}"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_MUTED
            p.space_after = Pt(4)

def add_flow_banner(slide, left, top, width, height, steps):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_ORANGE_LIGHT
    shape.line.color.rgb = COLOR_ORANGE

    n = len(steps)
    w = (width - 0.4 - (0.25 * (n - 1))) / n

    for i, s in enumerate(steps):
        s_left = left + 0.2 + i * (w + 0.25)
        node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(s_left), Inches(top + 0.12), Inches(w), Inches(height - 0.24))
        is_hl = "CKD" in s or "Best" in s or "Prediction" in s
        node.fill.solid()
        node.fill.fore_color.rgb = COLOR_ORANGE if is_hl else COLOR_WHITE
        node.line.color.rgb = COLOR_ORANGE if is_hl else COLOR_BORDER

        tf = node.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE if is_hl else COLOR_DARK
        p.alignment = PP_ALIGN.CENTER

        if i < n - 1:
            tb_a = slide.shapes.add_textbox(Inches(s_left + w), Inches(top + 0.12), Inches(0.25), Inches(height - 0.24))
            p_a = tb_a.text_frame.paragraphs[0]
            p_a.text = "➔"
            p_a.font.size = Pt(11)
            p_a.font.bold = True
            p_a.font.color.rgb = COLOR_ORANGE
            p_a.alignment = PP_ALIGN.CENTER

# Path to generated images
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_DASHBOARD = os.path.join(BASE_DIR, "images", "dashboard.jpg")
IMG_SHAP = os.path.join(BASE_DIR, "images", "shap_result.jpg")

# ==========================================
# SLIDE 1: PROJECT INTRODUCTION
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
set_white_bg(slide1)
add_header(slide1, "College Project Review 1 — Prototype Review", 
           "Explainable Multi-Model Machine Learning Framework for Early CKD Prediction", 
           '"An ML-based system for predicting the likelihood of Chronic Kidney Disease from patient health parameters"')

add_card(slide1, 0.8, 1.8, 5.7, 1.4, "Project Overview", [
    "Project: CKD Predict Prototype",
    "Domain: Healthcare + Machine Learning + Explainable AI",
    "Tech: React 19 + FastAPI + Python 3.14 + scikit-learn + SHAP"
])

if os.path.exists(IMG_DASHBOARD):
    slide1.shapes.add_picture(IMG_DASHBOARD, Inches(6.8), Inches(1.8), width=Inches(5.7))

add_flow_banner(slide1, 0.8, 5.0, 11.7, 1.0, ["Patient Health Data", "Data Preprocessing & Scaling", "Multi-Model Evaluation", "CKD Risk Prediction", "Explainable Result"])

# Disclaimer
disc = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.7), Inches(0.5))
disc.fill.solid()
disc.fill.fore_color.rgb = RGBColor(0xFF, 0xFB, 0xEB)
disc.line.color.rgb = RGBColor(0xFD, 0xE6, 0x8A)
p_d = disc.text_frame.paragraphs[0]
p_d.text = "⚠️ Research Disclaimer: Educational and research screening prototype — not a medical diagnosis tool."
p_d.font.size = Pt(11)
p_d.font.bold = True
p_d.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)

# ==========================================
# SLIDE 2: PROBLEM & OBJECTIVE
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
set_white_bg(slide2)
add_header(slide2, "Motivation & Vision", "Problem & Objective", "Addressing early detection challenges in renal healthcare through machine learning decision support")

add_card(slide2, 0.8, 1.8, 5.7, 3.2, "PROBLEM", [
    "Chronic Kidney Disease progresses gradually and can be hard to spot early.",
    "Manual clinical evaluation across 24+ lab parameters takes considerable time.",
    "Traditional ML models are black boxes, creating clinician skepticism."
], border_color=COLOR_RED)

add_card(slide2, 6.8, 1.8, 5.7, 3.2, "OBJECTIVE", [
    "Accept 24 patient clinical health parameters via a simple web interface.",
    "Automate data scrubbing, median imputation, and StandardScaler scaling.",
    "Evaluate 8 machine learning models with 5-Fold Stratified Cross-Validation.",
    "Select the best model, predict CKD likelihood, and display SHAP explanations."
], border_color=COLOR_ORANGE)

add_flow_banner(slide2, 0.8, 5.3, 11.7, 1.1, ["Patient Parameters", "Data Preprocessing", "ML Models", "Best Model", "CKD Prediction", "Explanation"])

# ==========================================
# SLIDE 3: DATASET & PATIENT PARAMETERS
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
set_white_bg(slide3)
add_header(slide3, "Data Architecture", "Dataset & Patient Parameters", "Empirical statistics from the UCI Chronic Kidney Disease benchmark dataset")

add_card(slide3, 0.8, 1.8, 3.6, 1.2, "400 Records", ["UCI CKD Benchmark Cohort", "Clinical Patient Instances"])
add_card(slide3, 4.8, 1.8, 3.6, 1.2, "24 Attributes", ["14 Numerical Features", "10 Nominal Categorical Features"])
add_card(slide3, 8.8, 1.8, 3.7, 1.2, "Target Distribution", ["250 CKD Instances (62.5%)", "150 Normal Instances (37.5%)"])

add_card(slide3, 0.8, 3.2, 3.6, 2.7, "PATIENT INFORMATION", [
    "Age (years)",
    "Blood Pressure (mmHg)"
])

add_card(slide3, 4.8, 3.2, 3.6, 2.7, "KIDNEY / BLOOD PARAMETERS", [
    "Specific Gravity, Albumin, Sugar",
    "Blood Glucose Random, Blood Urea",
    "Serum Creatinine, Sodium, Potassium",
    "Hemoglobin, PCV, WBC, RBC"
])

add_card(slide3, 8.8, 3.2, 3.7, 2.7, "MEDICAL HISTORY", [
    "Hypertension (yes / no)",
    "Diabetes Mellitus (yes / no)",
    "Coronary Artery Disease (yes / no)",
    "Appetite, Pedal Edema, Anemia"
])

add_flow_banner(slide3, 0.8, 6.1, 11.7, 0.9, ["Raw Dataset", "Cleaning", "Missing Value Handling", "Encoding", "Ready for ML"])

# ==========================================
# SLIDE 4: PROPOSED ML METHODOLOGY
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
set_white_bg(slide4)
add_header(slide4, "Core Architecture", "Proposed Methodology", "Multi-stage pipeline strictly fitted on training folds to prevent data leakage")

add_flow_banner(slide4, 0.8, 1.8, 11.7, 1.0, ["DATASET", "DATA PREPROCESSING", "FEATURE ENGINEERING", "FEATURE SELECTION", "CLASS BALANCING", "MODEL TRAINING", "MODEL COMPARISON", "BEST MODEL", "PREDICTION", "EXPLAINABLE AI"])

add_card(slide4, 0.8, 3.0, 3.6, 2.8, "Preprocessing [Implemented]", [
    "• Missing value handling (Median/Mode)",
    "• One-Hot Categorical Encoding",
    "• StandardScaler Normalization",
    "• Train/Test Split"
])

add_card(slide4, 4.8, 3.0, 3.6, 2.8, "Feature Selection [Implemented]", [
    "• Pearson Correlation Analysis",
    "• Chi-Square (χ²) Test",
    "• Recursive Feature Elimination (RFE)",
    "• LASSO L1 & Mutual Information"
])

add_card(slide4, 8.8, 3.0, 3.7, 2.8, "Class Balancing [Implemented]", [
    "• SMOTE (Synthetic Over-sampling)",
    "• SMOTETomek Hybrid Oversampling",
    "• Stratified 5-Fold Cross Validation",
    "• Strict zero data leakage"
])

add_card(slide4, 0.8, 6.0, 11.7, 1.0, "Evaluated Classifiers [8 Models Implemented]", [
    "Logistic Regression | Decision Tree | Random Forest | SVM | KNN | Naive Bayes (Best - 100% CV Acc) | XGBoost | Neural Network (MLP)"
])

# ==========================================
# SLIDE 5: PROTOTYPE WORKFLOW & UI
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
set_white_bg(slide5)
add_header(slide5, "User Interface", "Prototype Workflow & UI Modules", "Architecture of the live React 19 + Vite + FastAPI web application")

add_card(slide5, 0.8, 1.8, 5.7, 5.0, "Prototype UI Modules [All Live & Verified]", [
    "1. Dashboard Page: Telemetry badges, quick action CTAs, recent SQLite audit table.",
    "2. CKD Prediction Form: 24 Parameters, 3 sections, High/Low Risk sample presets.",
    "3. Prediction Result Page: High/Low Risk Tier badge, model probability % score.",
    "4. Explainable AI Explorer: SHAP feature waterfall chart & LIME rule weight table.",
    "5. Research Paper Generator: Auto-generates 15-section paper with PDF export."
])

if os.path.exists(IMG_SHAP):
    slide5.shapes.add_picture(IMG_SHAP, Inches(6.8), Inches(1.8), width=Inches(5.7))

# ==========================================
# SLIDE 6: EXPECTED OUTPUT & FUTURE WORK
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
set_white_bg(slide6)
add_header(slide6, "Deliverables & Future Roadmap", "Expected Output & Future Work", "Summary of completed prototype deliverables vs planned research enhancements")

add_card(slide6, 0.8, 1.8, 5.7, 3.9, "EXPECTED OUTPUT [✓ Completed & Tested]", [
    "✓ CKD / Not CKD prediction endpoint",
    "✓ Model-estimated probability percentage",
    "✓ Multi-model comparison across 8 classifiers",
    "✓ Best model identification (Naive Bayes 100% Acc)",
    "✓ Confusion matrix & Accuracy/F1/ROC-AUC metrics",
    "✓ SHAP & LIME Explainable AI attributions",
    "✓ Prediction history & SQLite audit log",
    "✓ Research report generation (15-section paper)"
], border_color=COLOR_GREEN)

add_card(slide6, 6.8, 1.8, 5.7, 3.9, "FUTURE WORK [Planned]", [
    "• Larger and more diverse clinical datasets",
    "• External prospective cohort validation",
    "• Clinical trial validation in outpatient settings",
    "• Improved model calibration over time",
    "• Longitudinal patient data tracking",
    "• Secure healthcare EHR / FHIR integration"
], border_color=RGBColor(0xF5, 0x9E, 0x0B))

add_flow_banner(slide6, 0.8, 5.9, 11.7, 0.9, ["Patient Data", "Preprocessing", "Multiple ML Models", "Best Model", "CKD Prediction", "Explainable AI", "Research Output"])

# Save Output
PPTX_PATH = os.path.join(BASE_DIR, "CKD_Predict_First_Review.pptx")
prs.save(PPTX_PATH)
print(f"Updated presentation saved at: {PPTX_PATH}")
